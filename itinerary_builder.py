import json
from datetime import datetime, date, timedelta
from collections import defaultdict
from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


# -------------------------------------------------------
#  Helper: Parse ISO date strings (handles Z, +01:00 etc.)
# -------------------------------------------------------
def parse_iso_date(dt_str):
    if not dt_str:
        return None
    if dt_str.endswith("Z"):
        dt_str = dt_str.replace("Z", "+00:00")
    return datetime.fromisoformat(dt_str).date()


# -------------------------------------------------------
#  Build per-day activities & accommodations
# -------------------------------------------------------
def build_day_structures(booking_json):
    attrs = booking_json["data"]["attributes"]
    start_date = date.fromisoformat(attrs["startDate"])
    end_date = date.fromisoformat(attrs["endDate"])

    included = booking_json.get("included", [])

    # activities by start day (non-accommodation)
    activities_by_day = defaultdict(list)
    # accommodations by every night: check-in <= day < check-out
    acc_by_day = defaultdict(list)

    for comp in included:
        c = comp.get("attributes", {})
        ctype = c.get("type")

        sday = parse_iso_date(c.get("startDateTime"))
        eday = parse_iso_date(c.get("endDateTime"))

        if ctype != "accommodation":
            if sday:
                activities_by_day[sday].append(c)
        else:
            # accommodation: occupies all nights check-in .. (check-out - 1)
            if sday and eday and sday <= eday:
                d = sday
                while d < eday:
                    acc_by_day[d].append(c)
                    d += timedelta(days=1)

    # Build ordered list of (day, activities_today, accommodations_today)
    day_entries = []
    d = start_date
    while d <= end_date:
        acts = activities_by_day.get(d, [])
        accs = acc_by_day.get(d, [])
        # Skip completely empty days (no travel, no hotel)
        if acts or accs:
            day_entries.append((d, acts, accs))
        d += timedelta(days=1)

    return day_entries


# -------------------------------------------------------
# Insert a paragraph *after* another paragraph
# -------------------------------------------------------
def insert_paragraph_after(paragraph, text="", style=None):
    new_p = OxmlElement("w:p")
    paragraph._element.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if text:
        new_para.add_run(text)
    if style:
        new_para.style = style
    return new_para


# -------------------------------------------------------
# Build the itinerary block dynamically
# -------------------------------------------------------
def fill_itinerary_block(doc, booking_json):
    placeholder = "{{ITINERARY_BLOCK}}"

    # 1) Find placeholder location
    anchor_para = None
    for p in doc.paragraphs:
        if placeholder in p.text:
            anchor_para = p
            break
    if anchor_para is None:
        raise ValueError("Placeholder {{ITINERARY_BLOCK}} not found in template")

    # 2) Build per-day structure (now includes leisure days)
    day_entries = build_day_structures(booking_json)

    current_para = anchor_para

    for day, activities, accommodations in day_entries:
        # --------------------------
        # DATE  (Heading 6)
        # --------------------------
        date_text = day.strftime("%A, %d %B %Y")
        current_para = insert_paragraph_after(
            current_para, date_text, style="Heading 6"
        )

        # --------------------------
        # ROUTE or LEISURE (Heading 3)
        # --------------------------
        main_activity = None
        for comp in activities:
            loc_start = comp.get("startLocation") or {}
            loc_end = comp.get("endLocation") or {}
            if loc_start.get("city") or loc_end.get("city"):
                main_activity = comp
                break
        if not main_activity and activities:
            main_activity = activities[0]

        # determine city from activity or accommodation
        city_from_activity = None
        if main_activity:
            loc_start = main_activity.get("startLocation") or {}
            loc_end = main_activity.get("endLocation") or {}
            city_from_activity = loc_start.get("city") or loc_end.get("city")

        city_from_acc = None
        if accommodations:
            loc = accommodations[0].get("startLocation") or {}
            city_from_acc = loc.get("city")

        city = city_from_activity or city_from_acc or "Switzerland"

        route_text = ""
        if main_activity:
            start_city = (main_activity.get("startLocation") or {}).get("city")
            end_city = (main_activity.get("endLocation") or {}).get("city")
            if start_city and end_city and start_city != end_city:
                route_text = f"{start_city} to {end_city}"
            elif start_city or end_city:
                route_text = start_city or end_city

        if route_text:
            heading3_text = route_text
        else:
            # no activity that day -> leisure day in that city
            heading3_text = f"Day at leisure in {city}"

        current_para = insert_paragraph_after(
            current_para, heading3_text, style="Heading 3"
        )

        # --------------------------
        # MAIN ACTIVITY TITLE (Normal)
        # --------------------------
        if main_activity:
            main_title = main_activity.get("productName") or main_activity.get("desc")
            if main_title:
                current_para = insert_paragraph_after(current_para, main_title)

        # --------------------------
        # SERVICES INCLUDED (only if activities)
        # --------------------------
        if activities:
            current_para = insert_paragraph_after(current_para, "Services Included:")
            for comp in activities:
                desc = comp.get("desc") or comp.get("productName") or ""
                bullet = desc
                if "," in desc:
                    bullet = desc.split(",", 1)[1].strip()
                current_para = insert_paragraph_after(
                    current_para, bullet, style="List Paragraph"
                )

        # blank line before accommodation block
        current_para = insert_paragraph_after(current_para, "")

        # --------------------------
        # ACCOMMODATION BLOCK
        # --------------------------
        if accommodations:
            # Header "Your Accommodation in XYZ" (Heading 4)
            acc_city = None
            for acc in accommodations:
                loc = acc.get("startLocation") or {}
                if loc.get("city"):
                    acc_city = loc["city"]
                    break
            if not acc_city:
                acc_city = city

            acc_header = f"Your Accommodation in {acc_city}"
            current_para = insert_paragraph_after(
                current_para, acc_header, style="Heading 4"
            )

            # Hotel line (Normal) – assume same hotel per city per night
            first_acc = accommodations[0]
            hotel_name = first_acc.get("productName") or first_acc.get("desc") or ""
            hotel_city = (first_acc.get("startLocation") or {}).get("city") or acc_city

            hotel_line = hotel_name
            if hotel_city:
                hotel_line += f", {hotel_city}"
            current_para = insert_paragraph_after(current_para, hotel_line)

            # Room summary lines: "2 × Single Room – 2 Travellers" etc.
            room_groups = defaultdict(lambda: {"rooms": 0, "pax": 0})
            for acc in accommodations:
                room = acc.get("room") or {}
                desc = room.get("desc") or "Room"
                num_pax = room.get("numPax") or 0
                room_groups[desc]["rooms"] += 1
                room_groups[desc]["pax"] += num_pax

            for desc, info in room_groups.items():
                line = f"{info['rooms']} ×  {desc} - {info['pax']} Travellers"
                current_para = insert_paragraph_after(current_para, line)

            # Booking Number line (from accommodation suppliers)
            booking_numbers = sorted(
                {
                    (acc.get("supplier") or {}).get("bookingReference")
                    for acc in accommodations
                    if (acc.get("supplier") or {}).get("bookingReference")
                }
            )
            if booking_numbers:
                booking_line = "Booking Number: " + ", ".join(booking_numbers)
                current_para = insert_paragraph_after(current_para, booking_line)

            # Blank line after each day block
            current_para = insert_paragraph_after(current_para, "")
        else:
            # No accommodation: small spacing
            current_para = insert_paragraph_after(current_para, "")

    # 3) Remove template placeholder
    anchor_para._element.getparent().remove(anchor_para._element)


# -------------------------------------------------------
# MAIN: for test purposes
# -------------------------------------------------------
def main():
    TEMPLATE = "TemplateNezasaAttributesV2.docx"
    JSON_FILE = "nezasa_booking.json"  # or your gxys6fk8bjpj.json
    OUTPUT = "output_itinerary_test.docx"

    with open(JSON_FILE, "r", encoding="utf-8") as f:
        data = json.load(f)

    doc = Document(TEMPLATE)
    fill_itinerary_block(doc, data)
    doc.save(OUTPUT)
    print("Created:", OUTPUT)


if __name__ == "__main__":
    main()
