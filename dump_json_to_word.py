import json
from pathlib import Path
from docx import Document


JSON_PATH = Path("nezasa_booking.json")
OUTPUT_DOCX = Path("json_structure_for_template.docx")


def walk(obj, prefix=""):
    """
    Recursively walk through a dict/list and yield (path, value) pairs
    for simple (non-dict, non-list) values.
    """
    if isinstance(obj, dict):
        for k, v in obj.items():
            new_prefix = f"{prefix}.{k}" if prefix else k
            yield from walk(v, new_prefix)
    elif isinstance(obj, list):
        for i, v in enumerate(obj):
            new_prefix = f"{prefix}[{i}]"
            yield from walk(v, new_prefix)
    else:
        # primitive value
        yield prefix, obj


def main():
    if not JSON_PATH.exists():
        raise FileNotFoundError(f"JSON file not found: {JSON_PATH}")

    with JSON_PATH.open("r", encoding="utf-8") as f:
        data = json.load(f)

    doc = Document()
    doc.add_heading("Nezasa JSON – Field Overview", level=1)
    doc.add_paragraph(
        "This document lists all JSON paths with example values. "
        "Use this to design your Word template and decide which fields to show."
    )

    for path, value in walk(data):
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(f"{path} = ").bold = True
        p.add_run(str(value))

    doc.save(OUTPUT_DOCX)
    print(f"Created: {OUTPUT_DOCX}")


if __name__ == "__main__":
    main()
