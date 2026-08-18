import json
from datetime import datetime
from collections import defaultdict
from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


# -------------------------------------------------------
#  Helper: Parse ISO date strings (handles Z, +01:00 etc.)
# -------------------------------------------------------
def parse_iso_date(dt_str):
    if not dt_str:
        return None
    if dt_str.endswith("Z"):
        dt_str = dt_str.replace("Z", "+00:00")
    return datetime.fromisoformat(dt_str).date()


# -------------------------------------------------------
#  Group all components by calendar date
# -------------------------------------------------------
def group_components_by_date(booking_json):
    day_map = defaultdict(list)

    for comp in booking_json.get("included", []):
        attrs = comp.get("attributes", {})
        day = parse_iso_date(attrs.get("startDateTime"))
        if not day:
            continue
        day_map[day].append(attrs)

    # Return dict sorted by day
    return dict(sorted(day_map.items(), key=lambda x: x[0]))


# -------------------------------------------------------
# Insert a paragraph *after* another paragraph
# -------------------------------------------------------
def insert_paragraph_after(paragraph, text="", style=None):
    new_p = OxmlElement("w:p")
    paragraph._element.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if text:
        new_para.add_run(text)
    if style:
        new_para.style = style
    return new_para


# -------------------------------------------------------
# Build the itinerary block dynamically
# -------------------------------------------------------
def fill_itinerary_block(doc, booking_json):
    placeholder = "{{ITINERARY_BLOCK}}"

    # 1) Find placeholder location
    anchor_para = None
    for p in doc.paragraphs:
        if placeholder in p.text:
            anchor_para = p
            break
    if anchor_para is None:
        raise ValueError("Placeholder {{ITINERARY_BLOCK}} not found in template")

    # 2) Build per-day structure
    days = group_components_by_date(booking_json)

    current_para = anchor_para

    for day, components in days.items():
        # --------------------------
        # DATE  (Heading 6)
        # --------------------------
        date_text = day.strftime("%A, %d %B %Y")
        current_para = insert_paragraph_after(
            current_para, date_text, style="Heading 6"
        )

        # Split components into activities & accommodations
        activities = [c for c in components if c.get("type") != "accommodation"]
        accommodations = [c for c in components if c.get("type") == "accommodation"]

        # --------------------------
        # ROUTE (Heading 3)
        # e.g. "Lucerne to Locarno"
        # --------------------------
        main_activity = None
        for comp in activities:
            loc_start = comp.get("startLocation") or {}
            loc_end = comp.get("endLocation") or {}
            if loc_start.get("city") or loc_end.get("city"):
                main_activity = comp
                break
        if not main_activity and activities:
            main_activity = activities[0]

        route_text = ""
        if main_activity:
            start_city = (main_activity.get("startLocation") or {}).get("city")
            end_city = (main_activity.get("endLocation") or {}).get("city")
            if start_city and end_city and start_city != end_city:
                route_text = f"{start_city} to {end_city}"
            elif start_city:
                route_text = start_city
            elif end_city:
                route_text = end_city

        if not route_text:
            # Fallback: use accommodation city or generic
            if accommodations:
                loc = accommodations[0].get("startLocation") or {}
                route_text = loc.get("city") or "Switzerland"
            else:
                route_text = "Switzerland"

        current_para = insert_paragraph_after(
            current_para, route_text, style="Heading 3"
        )

        # --------------------------
        # MAIN ACTIVITY TITLE (Normal)
        # e.g. "Gotthard Panorama Express Lucerne - Bellinzona - Locarno"
        # --------------------------
        if main_activity:
            main_title = main_activity.get("productName") or main_activity.get("desc")
            if main_title:
                current_para = insert_paragraph_after(current_para, main_title)

        # --------------------------
        # SERVICES INCLUDED (Normal + bullet lines)
        # --------------------------
        if activities:
            current_para = insert_paragraph_after(current_para, "Services Included:")
            for comp in activities:
                desc = comp.get("desc") or comp.get("productName") or ""
                # For things like
                # "Gotthard Panorama Express Lucerne - Bellinzona - Locarno, Seat reservation"
                # we only use the part after the comma for the bullet:
                bullet = desc
                if "," in desc:
                    bullet = desc.split(",", 1)[1].strip()

                current_para = insert_paragraph_after(
                    current_para, bullet, style="List Paragraph"
                )

        # blank line before accommodation block
        current_para = insert_paragraph_after(current_para, "")

        # --------------------------
        # ACCOMMODATION BLOCK
        # --------------------------
        if accommodations:
            # Header "Your Accommodation in XYZ" (Heading 4)
            acc_city = None
            for acc in accommodations:
                loc = acc.get("startLocation") or {}
                if loc.get("city"):
                    acc_city = loc["city"]
                    break
            if not acc_city:
                acc_city = route_text.split(" to ")[-1] if " to " in route_text else route_text

            acc_header = f"Your Accommodation in {acc_city}"
            current_para = insert_paragraph_after(
                current_para, acc_header, style="Heading 4"
            )

            # Hotel line (Normal)
            first_acc = accommodations[0]
            hotel_name = first_acc.get("productName") or first_acc.get("desc") or ""
            hotel_city = (first_acc.get("startLocation") or {}).get("city") or acc_city

            hotel_line = hotel_name
            if hotel_city:
                hotel_line += f" , {hotel_city}"
            current_para = insert_paragraph_after(current_para, hotel_line)

            # Summary line(s): "3 × Double Room Active - 4 Travellers"
            room_groups = defaultdict(lambda: {"rooms": 0, "pax": 0})
            for acc in accommodations:
                room = acc.get("room") or {}
                desc = room.get("desc") or "Room"
                num_pax = room.get("numPax") or 0
                room_groups[desc]["rooms"] += 1
                room_groups[desc]["pax"] += num_pax

            for desc, info in room_groups.items():
                line = f"{info['rooms']} ×  {desc} - {info['pax']} Travellers"
                current_para = insert_paragraph_after(current_para, line)

            # Booking Number line
            booking_numbers = sorted(
                {
                    (acc.get("supplier") or {}).get("bookingReference")
                    for acc in accommodations
                    if (acc.get("supplier") or {}).get("bookingReference")
                }
            )
            if booking_numbers:
                booking_line = "Booking Number: " + ", ".join(booking_numbers)
                current_para = insert_paragraph_after(current_para, booking_line)

            # Blank line after each day block
            current_para = insert_paragraph_after(current_para, "")
        else:
            # No accommodation: still add a bit of spacing
            current_para = insert_paragraph_after(current_para, "")

    # 3) Remove template placeholder
    anchor_para._element.getparent().remove(anchor_para._element)


# -------------------------------------------------------
# MAIN: for test purposes
# -------------------------------------------------------
def main():
    TEMPLATE = "TemplateNezasaAttributesV2.docx"
    JSON_FILE = "nezasa_booking.json"  # or your gxys6fk8bjpj.json for testing
    OUTPUT = "output_itinerary_test.docx"

    # Load JSON
    with open(JSON_FILE, "r", encoding="utf-8") as f:
        data = json.load(f)

    # Load Word template
    doc = Document(TEMPLATE)

    # Insert itinerary
    fill_itinerary_block(doc, data)

    # Save result
    doc.save(OUTPUT)
    print("Created:", OUTPUT)


if __name__ == "__main__":
    main()
